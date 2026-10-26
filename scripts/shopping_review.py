# -*- coding: utf-8 -*-
"""Googleショッピング改善案の確認用 Excel（除外キーワード＋商品名の書き換え一覧）。

使い方: python scripts/shopping_review.py <検索語の集計JSON> [出力パス]
  検索語の集計JSON = {検索語: [表示, クリック, 費用, 購入]}（PLA_v2 の search_term_view を集計したもの）
商品名は reports/google_shopping_titles.csv（ads_manager/shopping_titles.py が作成）を読む。
"""
from __future__ import annotations

import csv
import json
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from ads_manager.catalog_attributes import load_feed  # noqa: E402
from ads_manager.shopping_titles import NEGATIVES, TITLES_CSV, brand_label  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
HEAD = PatternFill("solid", fgColor="1F3864")
WHITE = Font(color="FFFFFF", bold=True)


def _header(ws, cols, widths):
    ws.append(cols)
    for i, w in enumerate(widths, 1):
        c = ws.cell(row=ws.max_row, column=i)
        c.fill, c.font = HEAD, WHITE
        ws.column_dimensions[c.column_letter].width = w
    ws.freeze_panes = f"A{ws.max_row + 1}"


def main():
    terms = json.load(open(sys.argv[1], encoding="utf-8"))
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "reports" / f"Googleショッピング_改善案_{date.today()}.xlsx"
    feed = load_feed()
    attrs = {r["id"]: r for r in csv.DictReader(open(ROOT / "reports" / "catalog_attributes.csv", encoding="utf-8"))}
    titles = [r for r in csv.DictReader(open(TITLES_CSV, encoding="utf-8"))
              if attrs.get(r["id"], {}).get("custom_label_0") == "regular" and feed.get(r["id"], {}).get("availability") == "in stock"]

    hit_terms = {k: v for k, v in terms.items() if any(w.lower() in k.lower() for w in NEGATIVES)}
    tot = [sum(v[i] for v in terms.values()) for i in (1, 2, 3)]
    neg = [sum(v[i] for v in hit_terms.values()) for i in (1, 2, 3)]

    wb = Workbook()
    ws = wb.active
    ws.title = "概要"
    lines = [
        ("Googleショッピング（PLA_v2）改善案", True),
        (f"作成日 {date.today()}　対象: 在庫ありの通常価格品 {len(titles)}点（アウトレット品は従来どおり広告しない）", False),
        ("", False),
        ("1. 除外キーワード（このシート「除外キーワード」）", True),
        ("　8/1〜10/25 の検索語で、購入につながっていない他社ブランド名・用途違いの語で検索した人には出さない。", False),
        ("　パタゴニア・アークテリクス・ホグロフスは購入が出ていたので除外しない。", False),
        (f"　8/1〜10/25 の実績: 除外する語を含む検索は クリック {neg[0]:,}回・費用 {neg[1]:,.0f}円・購入 {neg[2]:.0f}件"
         f"（全体 クリック {tot[0]:,}回・費用 {tot[1]:,.0f}円・購入 {tot[2]:.1f}件のうち。語どうしが重なる検索は1回だけ数えた）。", False),
        ("2. 商品名の書き換え（このシート「商品名」）", True),
        ("　サイトの商品データは英語の型番だけなので、Google 上の商品名を「ブランド＋元の商品名｜性別・種類」にする。", False),
        ("　あわせて Google の商品カテゴリ・性別を登録する。サイトの表示は変わらない。毎週月曜の定例で自動更新する。", False),
        ("　「確認」欄で 外したい／直したい を選び、コメントに直し方を書いてください。", False),
    ]
    for text, bold in lines:
        ws.append([text])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=bold, size=13 if ws.max_row == 1 else 11)
    ws.column_dimensions["A"].width = 120

    ws2 = wb.create_sheet("除外キーワード")
    _header(ws2, ["除外する語", "8/1〜10/25 クリック", "費用（円）", "購入"], [26, 18, 14, 10])
    for w in NEGATIVES:
        wl = w.lower()
        hit = [v for k, v in terms.items() if wl in k.lower()]
        ws2.append([w, sum(v[1] for v in hit), round(sum(v[2] for v in hit)), round(sum(v[3] for v in hit), 1)])

    ws3 = wb.create_sheet("商品名")
    _header(ws3, ["ブランド", "いまの商品名（Google 上）", "書き換え後", "種類の大分類", "性別", "Google 商品カテゴリ",
                  "商品ID", "商品ページ", "確認", "コメント"], [16, 40, 58, 12, 9, 46, 13, 12, 12, 30])
    dv = DataValidation(type="list", formula1='"OK,外したい,直したい"', allow_blank=True)
    ws3.add_data_validation(dv)
    for r in sorted(titles, key=lambda x: (brand_label(x["id"]), x["title"])):
        url = f"https://www.fullmarksstore.jp/item/{r['id']}.html"
        g = {"male": "メンズ", "female": "レディース"}.get(r["gender"], "キッズ" if r["age_group"] == "kids" else "男女兼用")
        ws3.append([brand_label(r["id"]), r["old_title"], r["title"], r["category"], g,
                    r["google_product_category"], r["id"], "開く", "", ""])
        row = ws3.max_row
        ws3.cell(row=row, column=8).hyperlink = url
        ws3.cell(row=row, column=8).font = Font(color="0563C1", underline="single")
        dv.add(ws3.cell(row=row, column=9))
        for col in (2, 3, 6):
            ws3.cell(row=row, column=col).alignment = Alignment(wrap_text=True, vertical="top")
    ws3.auto_filter.ref = ws3.dimensions
    wb.save(out)
    print(out)


if __name__ == "__main__":
    main()
