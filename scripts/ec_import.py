# -*- coding: utf-8 -*-
"""ユーザーから毎月もらう EC 実績の Excel（予実資料）から、月次レポート用の数値を copy/ec/<月>.json に書き出す。

使い方:
  python scripts/ec_import.py <予実資料のフォルダ または zip> --month 2026-09

読むファイル（ファイル名の一部で探す）:
  ・「WEB売上実績」 … 月別の売上（税抜）・会社予算・前年実績、ブランド別の月別売上
  ・「消化率」       … 新作（今シーズン品）のブランド別 目標と実績
  ・「売上内訳」     … 月ごとのシートに、ブランド別の 新作／旧品 の売上
EC全体の売上は API で取れないため、このファイルが唯一の情報源。すべて税抜。
"""
from __future__ import annotations

import argparse
import json
import re
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
FY_MONTHS = [5, 6, 7, 8, 9, 10, 11, 12, 1, 2, 3, 4]  # 年度は5月始まり
BRAND_NAME = {"NORRONA": "NORRØNA", "SR": "SAIL RACING", "PU": "PERMANENT UNION",
              "Permanent Union": "PERMANENT UNION", "FRANK D": "FRANK DANDY"}


def _find(files: list[Path], key: str) -> Path:
    for f in files:
        if key in f.name:
            return f
    raise SystemExit(f"「{key}」を含む xlsx が見つかりません")


def _files(src: Path) -> list[Path]:
    if src.suffix == ".zip":
        tmp = Path(tempfile.mkdtemp())
        with zipfile.ZipFile(src) as z:
            for info in z.infolist():
                name = info.filename
                if not info.flag_bits & 0x800:  # Windows で作った zip はファイル名が cp932
                    name = name.encode("cp437").decode("cp932", errors="replace")
                if name.endswith(".xlsx"):
                    (tmp / Path(name).name).write_bytes(z.read(info))
        src = tmp
    return sorted(src.rglob("*.xlsx"))


def _num(v) -> int:
    return round(float(v)) if isinstance(v, (int, float)) else 0


def web_sales(path: Path, month: int) -> dict:
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    col = 2 + FY_MONTHS.index(month)  # C列=5月
    by_label = {r[1]: r for r in rows if len(r) > col and isinstance(r[1], str)}
    act = by_label["2026年度実績"]; bud = by_label["2026年度予算"]; ly = by_label["2025年度実績"]
    upto = range(2, col + 1)
    out = {"sales": _num(act[col]), "budget": _num(bud[col]), "last_year": _num(ly[col]),
           "ytd": {"sales": sum(_num(act[c]) for c in upto), "budget": sum(_num(bud[c]) for c in upto),
                   "last_year": sum(_num(ly[c]) for c in upto)}}
    if col + 1 <= 13:
        out["next_month"] = {"budget": _num(bud[col + 1]), "last_year": _num(ly[col + 1])}
    brands, on = {}, False
    for r in rows:
        if r[1] == "2026年度" and not on:
            on = True; continue
        if on:
            if r[1] in ("月別TOTAL", None):
                break
            if r[1] != "その他":
                brands[BRAND_NAME.get(r[1], r[1].upper())] = {"sales": _num(r[col]), "prev_month": _num(r[col - 1])}
    out["brands"] = brands
    return out


def new_season(path: Path, year: int, month: int) -> dict:
    ws = openpyxl.load_workbook(path, data_only=True)["実績"]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    head = next(r for r in rows if r[0] == "WEBSTORE")
    names = [BRAND_NAME.get(h, h) for h in head]
    tgt = nxt = None
    for i, r in enumerate(rows):
        if isinstance(r[0], datetime) and r[0].month == month and r[1] == "売上目標":
            if tgt is None or r[0].year == year:
                tgt, act = r, rows[i + 1]
        if isinstance(r[0], datetime) and r[0].month == month % 12 + 1 and r[1] == "売上目標":
            nxt = r
    out = {"brands": {}}
    for j, n in enumerate(names):
        if j < 2 or n is None:
            continue
        if n == "合計":
            out["target"], out["sales"] = _num(tgt[j]), _num(act[j])
            if nxt:
                out["next_target"] = _num(nxt[j])
        else:
            out["brands"][n] = {"target": _num(tgt[j]), "sales": _num(act[j]),
                                "next_target": _num(nxt[j]) if nxt else 0.0}
    return out


def breakdown(path: Path, month: int) -> dict:
    ws = openpyxl.load_workbook(path, data_only=True)[f"{month}月"]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    hi = next(i for i, r in enumerate(rows) if "消化率対象売上" in r)
    head, rows = rows[hi], rows[hi:]
    c_new, c_old, c_tot = head.index("消化率対象売上"), head.index("旧品売上"), head.index("合計（税抜）")
    season = head[c_new - 1] or ""
    out = {"season": re.sub(r"\s+", "", season), "brands": {}}
    for r in rows[1:]:
        if r[1] is None:
            if r[c_tot] is not None and "total" not in out:
                out["total"] = {"new": _num(r[c_new]), "old": _num(r[c_old]), "sales": _num(r[c_tot])}
            continue
        b = BRAND_NAME.get(r[1], r[1].upper())
        out["brands"][b] = {"new": _num(r[c_new]), "old": _num(r[c_old]), "sales": _num(r[c_tot])}
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("src", help="予実資料のフォルダ または zip")
    ap.add_argument("--month", required=True, help="YYYY-MM")
    args = ap.parse_args()
    y, m = map(int, args.month.split("-"))
    files = _files(Path(args.src))
    web = web_sales(_find(files, "WEB売上実績"), m)
    ns = new_season(_find(files, "消化率"), y, m)
    bd = breakdown(_find(files, "売上内訳"), m)
    brands = []
    for b, v in sorted(web["brands"].items(), key=lambda kv: -kv[1]["sales"]):
        if not v["sales"] and not ns["brands"].get(b, {}).get("target"):
            continue
        n = ns["brands"].get(b, {}); d = bd["brands"].get(b, {})
        brands.append({"brand": b, "sales": v["sales"], "prev_month": v["prev_month"],
                       "new": d.get("new", n.get("sales", 0.0)), "old": d.get("old", 0.0),
                       "new_target": n.get("target", 0.0), "next_new_target": n.get("next_target", 0.0)})
    out = {"_comment": "scripts/ec_import.py が予実資料（ユーザー提供の Excel）から生成。すべて税抜。"
                       "新作=今シーズン品（会社の消化率の対象）、旧品=前シーズン以前（セール品の大半）",
           "month": args.month, "sales": web["sales"], "budget": web["budget"], "last_year": web["last_year"],
           "ytd": web["ytd"], "next_month": {**web.get("next_month", {}), "new_target": ns.get("next_target", 0.0)},
           "new_season": {"label": bd["season"], "sales": bd.get("total", {}).get("new", ns["sales"]),
                          "target": ns["target"]},
           "old_items": bd.get("total", {}).get("old", 0.0), "brands": brands}
    p = ROOT / "copy" / "ec" / f"{args.month}.json"
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"書き出し: {p}")
    print(f"売上 {out['sales']:,.0f}（予算 {out['budget']:,.0f} / 前年 {out['last_year']:,.0f}）"
          f" 新作 {out['new_season']['sales']:,.0f}／目標 {out['new_season']['target']:,.0f}")


if __name__ == "__main__":
    main()
