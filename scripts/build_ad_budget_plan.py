# -*- coding: utf-8 -*-
"""広告予算計画（2026-05〜2027-04、税抜、年間上限1,200万）を数式ベースのExcelで生成する。

使い方: python scripts/build_ad_budget_plan.py [出力パス]
シート: 前提 / 月別予算 / FULLMARKS内訳 / 実施_日予算 / 9月実績
"""
from __future__ import annotations

import calendar
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

F = "Arial"; YEN = "¥#,##0;(¥#,##0);-"; PCT = "0.0%"
blue = Font(name=F, size=10, color="0000FF"); black = Font(name=F, size=10)
bold = Font(name=F, size=10, bold=True); green = Font(name=F, size=10, color="008000")
wb_ = Font(name=F, size=10, bold=True, color="FFFFFF"); small = Font(name=F, size=8, italic=True)
yellow = PatternFill("solid", fgColor="FFFF00"); hdr = PatternFill("solid", fgColor="1F3864")
sub = PatternFill("solid", fgColor="D9E1F2"); grey = PatternFill("solid", fgColor="EDEDED")
newf = PatternFill("solid", fgColor="E2EFDA")
thin = Side(style="thin", color="BFBFBF"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)

ANNUAL = 12_000_000
ACTUALS = [("2026-05", 280_564, 0, "FULLMARKS API実績"),
           ("2026-06", 298_599, 651_401, "月合計95万（ユーザー提供）。差額はMDX管理口座等"),
           ("2026-07", 328_194, 621_806, "月合計95万（ユーザー提供）"),
           ("2026-08", 468_641, 0, "FULLMARKS API実績。他ストア未稼働")]
SEP_MTD = 220_464          # 9/1〜9/10 実績（Google 104,706 + Meta 115,758）
SEP_MTD_DAYS = 10
SEP_NORMAL = 650_000       # 9月 通常運用（FULLMARKSのみ）
EVENTS = {"2026-09": ("シルバーウィーク企画", 100_000), "2026-12": ("年末年始企画", 100_000),
          "2027-02": ("冬セール", 250_000), "2027-03": ("イベント告知予備", 50_000)}
MONTHS_OCT = ["2026-10", "2026-11", "2026-12", "2027-01", "2027-02", "2027-03", "2027-04"]
WEIGHTS = [0.15, 0.155, 0.17, 0.14, 0.16, 0.115, 0.11]
STORE_R = 0.15
BRANDS = [  # (名称, 9月比率, 10月〜比率, 備考)
    ("HOUDINI", 0.27, 0.20, "売上の7割。Google指名 ROAS 12〜13、Meta RTG 5.0（9月第1週）。HOUDINI STORE稼働後は縮小"),
    ("商品連動(ショッピング/カタログ)", 0.30, 0.28, "通常価格品のみ（アウトレットは広告しない）。PLA 4,000円/日＋Metaカタログ9本 5,000円/日"),
    ("店舗指名(フルマークス)", 0.17, 0.18, "ROAS 9〜10の安定枠"),
    ("POC", 0.09, 0.12, "9月売上構成比13%に上昇。専用ストア無し"),
    ("NORRØNA", 0.07, 0.06, "ROAS 2前後。NORRONA STORE移管で最小化"),
    ("ACLIMA", 0.05, 0.06, "ユーザー判断で維持（売れる見込み）"),
    ("HESTRA【新規】", 0.04, 0.08, "10〜1月が本番。フィード未掲載の解消が前提"),
    ("KANG【新規】", 0.005, 0.01, "テスト枠"),
    ("PLUS ONE WORKS【新規】", 0.005, 0.01, "テスト枠"),
]
CAMP = {  # ブランド -> [(媒体, キャンペーン, 比率, 新設?)]
    "HOUDINI": [("Google", "UC_SK_1_指名_フーディ二", 0.65, False), ("Meta", "UC_DN_3_RTG_フーディニ_2608", 0.35, False)],
    "商品連動(ショッピング/カタログ)": [("Google", "UC_PL_5_PLA_v2（通常価格のみ）", 0.45, False), ("Meta", "UC_DN_3_CVS_フルマークス（ブランド別9本）", 0.55, False)],
    "店舗指名(フルマークス)": [("Google", "UC_SK_1_指名_フルマークス", 1.0, False)],
    "POC": [("Google", "UC_SK_1_指名_ポック", 0.45, False), ("Meta", "UC_DN_3_RTG_ポック_2608", 0.55, False)],
    "NORRØNA": [("Google", "UC_SK_1_指名_ノローナ", 0.60, False), ("Meta", "UC_DN_3_RTG_ノローナ_2608", 0.40, False)],
    "ACLIMA": [("Meta", "UC_DN_4_CLK_アクリマ", 0.60, False), ("Meta", "UC_DN_3_RTG_アクリマ_2608", 0.40, False)],
    "HESTRA【新規】": [("Google", "【新設】UC_SK_1_指名_HESTRA", 1.0, True)],
    "KANG【新規】": [("Google", "【新設】UC_SK_1_指名_KANG", 1.0, True)],
    "PLUS ONE WORKS【新規】": [("Google", "【新設】UC_SK_1_指名_PLUS ONE WORKS", 1.0, True)],
}
STORE_CAMP = [("Google", "【新設】指名検索", 0.45), ("Google", "【新設】ショッピング(通常価格のみ)", 0.30), ("Meta", "【新設】カタログ/リターゲティング", 0.25)]
STORES = ["HOUDINI STORE", "NORRONA STORE", "PU STORE"]
SEP_ACTUAL_BY_CAMP = [("Google", "UC_PL_5_PLA_v2", 42_548), ("Google", "UC_SK_1_指名_フーディ二", 21_682),
                      ("Google", "UC_SK_1_指名_フルマークス", 20_073), ("Google", "UC_SK_1_指名_ノローナ", 17_816),
                      ("Google", "UC_SK_1_指名_ポック", 2_540), ("Meta", "UC_DN_3_CVS_フルマークス", 42_212),
                      ("Meta", "UC_DN_3_RTG_フーディニ_2608", 22_086), ("Meta", "UC_DN_3_RTG_アクリマ_2608", 14_179),
                      ("Meta", "UC_DN_3_RTG_ポック_2608", 12_428), ("Meta", "UC_DN_3_RTG_ノローナ_2608", 12_315),
                      ("Meta", "UC_DN_4_CLK_アクリマ", 12_149), ("Google/Meta", "アウトレット枠（9/7停止）", 436)]


def days(m: str) -> int:
    return calendar.monthrange(int(m[:4]), int(m[5:]))[1]


def build(out: str) -> None:
    wb = Workbook()
    # ================= 前提 =================
    ws = wb.active; ws.title = "前提"
    ws["A1"] = "広告予算計画 2026-05〜2027-04（税抜）― 前提・入力値　2026-09-11 改訂"; ws["A1"].font = Font(name=F, size=13, bold=True)
    ws["A2"] = ("青字＝入力値、黄色＝仮置き（他ストアの実績や方針で差し替え）。他シートは全てこのシートを参照。"
                "方針: アウトレット品は広告しない／広告口座は FULLMARKS広告運用(Meta)・FULLMARKS Inc.(Google) のみ。"); ws["A2"].font = small
    r = 4
    ws.cell(row=r, column=1, value="年間広告費上限（4ストア合計・税抜）").font = black
    c = ws.cell(row=r, column=2, value=ANNUAL); c.font = blue; c.number_format = YEN
    K_ANNUAL = "前提!$B$4"
    ws.cell(row=5, column=1, value="3ストア各社の予算シェア（10月〜、各）").font = black
    c = ws.cell(row=5, column=2, value=STORE_R); c.font = blue; c.number_format = PCT; c.fill = yellow
    K_STORE = "前提!$B$5"
    ws.cell(row=5, column=3, value="HOUDINI STORE 9月末開店 / NORRONA・PU 10月開店 → 10月から各15%、FULLMARKS 55%").font = small
    ws.cell(row=6, column=1, value="9月 通常運用（FULLMARKS）").font = black
    c = ws.cell(row=6, column=2, value=SEP_NORMAL); c.font = blue; c.number_format = YEN
    K_SEP = "前提!$B$6"
    ws.cell(row=6, column=3, value=f"9/1〜9/10実績 ¥{SEP_MTD:,}（日割 ¥{SEP_MTD//SEP_MTD_DAYS:,}、月換算 ¥{SEP_MTD*30//SEP_MTD_DAYS:,}）→ 65万ペースに乗っている").font = small
    # 実績
    ws["A9"] = "実績（5〜8月）"; ws["A9"].font = bold
    for c_, h in enumerate(["月", "FULLMARKS", "その他(MDX等)", "合計", "備考"], 1):
        cell = ws.cell(row=10, column=c_, value=h); cell.font = wb_; cell.fill = hdr; cell.alignment = center
    for i, (m, fm, other, note) in enumerate(ACTUALS, 11):
        ws.cell(row=i, column=1, value=m).font = black
        for c_, v in ((2, fm), (3, other)):
            cell = ws.cell(row=i, column=c_, value=v); cell.font = blue; cell.number_format = YEN
        ws.cell(row=i, column=4, value=f"=B{i}+C{i}").number_format = YEN
        ws.cell(row=i, column=5, value=note).font = small
    ws.cell(row=15, column=1, value="5〜8月 計").font = bold
    for c_ in (2, 3, 4):
        cell = ws.cell(row=15, column=c_, value=f"=SUM({L(c_)}11:{L(c_)}14)"); cell.font = bold; cell.number_format = YEN
    K_ACT = "前提!$D$15"
    # イベント
    ws["A18"] = "イベント予備費（通常運用と別枠）"; ws["A18"].font = bold
    for c_, h in enumerate(["月", "内容", "金額"], 1):
        cell = ws.cell(row=19, column=c_, value=h); cell.font = wb_; cell.fill = hdr; cell.alignment = center
    ev_rows = {}
    for i, (m, (name, amt)) in enumerate(EVENTS.items(), 20):
        ws.cell(row=i, column=1, value=m).font = black; ws.cell(row=i, column=2, value=name).font = black
        cell = ws.cell(row=i, column=3, value=amt); cell.font = blue; cell.number_format = YEN; cell.fill = yellow
        ev_rows[m] = i
    ev_end = 19 + len(EVENTS)
    ws.cell(row=ev_end + 1, column=1, value="計").font = bold
    cell = ws.cell(row=ev_end + 1, column=3, value=f"=SUM(C20:C{ev_end})"); cell.font = bold; cell.number_format = YEN
    K_EV = f"前提!$C${ev_end+1}"
    # 10〜4月の月別比率
    r0 = ev_end + 4
    ws.cell(row=r0 - 1, column=1, value="10〜4月 通常運用の月別配分比率（繁忙期に厚く。4月は残額調整）").font = bold
    for c_, h in enumerate(["月", "比率"], 1):
        cell = ws.cell(row=r0, column=c_, value=h); cell.font = wb_; cell.fill = hdr; cell.alignment = center
    w_rows = {}
    for i, (m, w) in enumerate(zip(MONTHS_OCT, WEIGHTS), r0 + 1):
        ws.cell(row=i, column=1, value=m).font = black
        cell = ws.cell(row=i, column=2, value=w); cell.font = blue; cell.number_format = PCT; cell.fill = yellow
        w_rows[m] = i
    wr_end = r0 + len(MONTHS_OCT)
    ws.cell(row=wr_end + 1, column=1, value="計（100%）").font = bold
    cell = ws.cell(row=wr_end + 1, column=2, value=f"=SUM(B{r0+1}:B{wr_end})"); cell.font = bold; cell.number_format = PCT
    ws.cell(row=wr_end + 2, column=1, value="10〜4月 通常運用の総額（上限−実績−9月−イベント）").font = bold
    cell = ws.cell(row=wr_end + 2, column=2, value=f"={K_ANNUAL}-{K_ACT}-{K_SEP}-{K_EV}"); cell.font = bold; cell.number_format = YEN
    K_ALLOC = f"前提!$B${wr_end+2}"
    # ブランド比率
    b0 = wr_end + 5
    ws.cell(row=b0 - 1, column=1, value="FULLMARKS 内のブランド配分比率").font = bold
    for c_, h in enumerate(["ブランド/枠", "9月 比率", "10月〜 比率", "備考"], 1):
        cell = ws.cell(row=b0, column=c_, value=h); cell.font = wb_; cell.fill = hdr; cell.alignment = center
    b_rows = {}
    for i, (name, r9, r10, note) in enumerate(BRANDS, b0 + 1):
        ws.cell(row=i, column=1, value=name).font = black
        for c_, v in ((2, r9), (3, r10)):
            cell = ws.cell(row=i, column=c_, value=v); cell.font = blue; cell.number_format = PCT; cell.fill = yellow
        ws.cell(row=i, column=4, value=note).font = small
        b_rows[name] = i
    b_end = b0 + len(BRANDS)
    ws.cell(row=b_end + 1, column=1, value="計（各100%）").font = bold
    for c_ in (2, 3):
        cell = ws.cell(row=b_end + 1, column=c_, value=f"=SUM({L(c_)}{b0+1}:{L(c_)}{b_end})"); cell.font = bold; cell.number_format = PCT
    for col, w in zip("ABCDE", (40, 14, 14, 14, 70)):
        ws.column_dimensions[col].width = w

    # ================= 月別予算 =================
    ws2 = wb.create_sheet("月別予算")
    ws2["A1"] = "月別 広告予算（税抜）― 4ストア合計で年間1,200万"; ws2["A1"].font = Font(name=F, size=13, bold=True)
    ws2["A2"] = "通常運用＝上限−実績−9月−イベント予備費 を10〜4月の比率で配分。10月から HOUDINI / NORRONA / PU STORE が各15%。"; ws2["A2"].font = small
    h = ["月", "区分", "FULLMARKS", "HOUDINI STORE", "NORRONA STORE", "PU STORE", "その他(実績)", "通常運用 計", "イベント予備費", "月合計", "累計", "上限までの残り", "備考"]
    for c_, hh in enumerate(h, 1):
        cell = ws2.cell(row=4, column=c_, value=hh); cell.font = wb_; cell.fill = hdr; cell.alignment = center
    row = 5; month_rows = {}
    for i, (m, fm, other, note) in enumerate(ACTUALS):
        ws2.cell(row=row, column=1, value=m); ws2.cell(row=row, column=2, value="実績")
        ws2.cell(row=row, column=3, value=f"=前提!B{11+i}").font = green
        for c_ in (4, 5, 6): ws2.cell(row=row, column=c_, value=0)
        ws2.cell(row=row, column=7, value=f"=前提!C{11+i}").font = green
        ws2.cell(row=row, column=8, value=f"=C{row}+D{row}+E{row}+F{row}+G{row}")
        ws2.cell(row=row, column=9, value=0)
        ws2.cell(row=row, column=10, value=f"=H{row}+I{row}")
        ws2.cell(row=row, column=11, value=f"=J{row}" if row == 5 else f"=K{row-1}+J{row}")
        ws2.cell(row=row, column=12, value=f"={K_ANNUAL}-K{row}")
        ws2.cell(row=row, column=13, value=note).font = small
        for c_ in range(1, 13): ws2.cell(row=row, column=c_).fill = grey
        month_rows[m] = row; row += 1
    # 9月
    ws2.cell(row=row, column=1, value="2026-09"); ws2.cell(row=row, column=2, value="計画")
    ws2.cell(row=row, column=3, value=f"={K_SEP}").font = green
    for c_ in (4, 5, 6, 7): ws2.cell(row=row, column=c_, value=0)
    ws2.cell(row=row, column=8, value=f"=C{row}+D{row}+E{row}+F{row}+G{row}")
    ws2.cell(row=row, column=9, value=f"=前提!C{ev_rows['2026-09']}").font = green
    ws2.cell(row=row, column=10, value=f"=H{row}+I{row}")
    ws2.cell(row=row, column=11, value=f"=K{row-1}+J{row}")
    ws2.cell(row=row, column=12, value=f"={K_ANNUAL}-K{row}")
    ws2.cell(row=row, column=13, value=f"FULLMARKSのみ。9/1〜10実績 ¥{SEP_MTD:,}（65万ペース）。SW企画10万は別枠").font = small
    month_rows["2026-09"] = row; row += 1
    for m in MONTHS_OCT:
        ws2.cell(row=row, column=1, value=m); ws2.cell(row=row, column=2, value="計画")
        ws2.cell(row=row, column=8, value=(f"={K_ALLOC}*前提!$B${w_rows[m]}" if m != "2027-04"
                                            else f"={K_ALLOC}-SUM(H{month_rows['2026-10']}:H{row-1})"))
        for c_ in (4, 5, 6): ws2.cell(row=row, column=c_, value=f"=ROUND(H{row}*{K_STORE},0)")
        ws2.cell(row=row, column=3, value=f"=H{row}-D{row}-E{row}-F{row}")
        ws2.cell(row=row, column=7, value=0)
        ws2.cell(row=row, column=9, value=(f"=前提!C{ev_rows[m]}" if m in ev_rows else 0))
        if m in ev_rows: ws2.cell(row=row, column=9).font = green
        ws2.cell(row=row, column=10, value=f"=H{row}+I{row}")
        ws2.cell(row=row, column=11, value=f"=K{row-1}+J{row}")
        ws2.cell(row=row, column=12, value=f"={K_ANNUAL}-K{row}")
        note = {"2026-10": "3ストア稼働・広告開始。勝負月", "2026-12": "年末年始企画 予備10万", "2027-02": "冬セール 予備25万",
                "2027-03": "イベント告知予備5万", "2027-04": "調整月（残額を全額割当）"}.get(m, "")
        ws2.cell(row=row, column=13, value=note).font = small
        month_rows[m] = row; row += 1
    tr = row
    ws2.cell(row=tr, column=1, value="年間合計").font = bold
    for c_ in range(3, 11):
        cell = ws2.cell(row=tr, column=c_, value=f"=SUM({L(c_)}5:{L(c_)}{tr-1})"); cell.font = bold; cell.fill = sub
    ws2.cell(row=tr, column=12, value=f"={K_ANNUAL}-J{tr}").font = bold
    ws2.cell(row=tr, column=13, value="上限までの残りが0であること").font = small
    for rr in range(5, tr + 1):
        for c_ in range(3, 13): ws2.cell(row=rr, column=c_).number_format = YEN
        for c_ in range(1, 14): ws2.cell(row=rr, column=c_).border = border
    for col, w in zip("ABCDEFGHIJKLM", (11, 7, 13, 14, 14, 12, 12, 13, 13, 13, 14, 14, 46)):
        ws2.column_dimensions[col].width = w
    PLAN_MONTHS = ["2026-09"] + MONTHS_OCT

    # ================= FULLMARKS内訳 =================
    ws3 = wb.create_sheet("FULLMARKS内訳")
    ws3["A1"] = "FULLMARKS STORE ブランド別予算（月別予算のFULLMARKS列 × 前提の比率）"; ws3["A1"].font = Font(name=F, size=13, bold=True)
    ws3["A2"] = "最終行は差引で列合計をFULLMARKS月額に一致させる。アウトレット品はどの枠でも広告しない。"; ws3["A2"].font = small
    for c_, hh in enumerate(["ブランド/枠"] + PLAN_MONTHS + ["年間(9〜4月)", "備考"], 1):
        cell = ws3.cell(row=4, column=c_, value=hh); cell.font = wb_; cell.fill = hdr; cell.alignment = center
    n = len(BRANDS); first = 5; last = first + n - 1
    for i, (name, r9, r10, note) in enumerate(BRANDS):
        rr = first + i
        ws3.cell(row=rr, column=1, value=name).font = black
        for mi, m in enumerate(PLAN_MONTHS):
            col = 2 + mi; ratio_col = "B" if m == "2026-09" else "C"
            src = f"月別予算!$C${month_rows[m]}"
            if i < n - 1:
                f = f"=ROUND({src}*前提!${ratio_col}${b_rows[name]},0)"
            else:
                f = f"={src}-SUM({L(col)}{first}:{L(col)}{rr-1})"
            cell = ws3.cell(row=rr, column=col, value=f); cell.font = green; cell.number_format = YEN
        cell = ws3.cell(row=rr, column=2 + len(PLAN_MONTHS), value=f"=SUM(B{rr}:{L(1+len(PLAN_MONTHS))}{rr})"); cell.font = bold; cell.number_format = YEN
        ws3.cell(row=rr, column=3 + len(PLAN_MONTHS), value=note).font = small
        if "新規" in name:
            for c_ in range(1, 3 + len(PLAN_MONTHS)): ws3.cell(row=rr, column=c_).fill = newf
    tr3 = last + 1
    ws3.cell(row=tr3, column=1, value="合計").font = bold
    for col in range(2, 3 + len(PLAN_MONTHS)):
        cell = ws3.cell(row=tr3, column=col, value=f"=SUM({L(col)}{first}:{L(col)}{last})"); cell.font = bold; cell.number_format = YEN; cell.fill = sub
    ws3.cell(row=tr3 + 1, column=1, value="月別予算との差（0）").font = bold
    for mi, m in enumerate(PLAN_MONTHS):
        col = 2 + mi
        cell = ws3.cell(row=tr3 + 1, column=col, value=f"={L(col)}{tr3}-月別予算!$C${month_rows[m]}"); cell.number_format = YEN
    for rr in range(4, tr3 + 2):
        for c_ in range(1, 4 + len(PLAN_MONTHS)): ws3.cell(row=rr, column=c_).border = border
    ws3.column_dimensions["A"].width = 30
    for col in range(2, 3 + len(PLAN_MONTHS)): ws3.column_dimensions[L(col)].width = 11
    ws3.column_dimensions[L(3 + len(PLAN_MONTHS))].width = 60

    # ================= 実施_日予算 =================
    ws4 = wb.create_sheet("実施_日予算")
    ws4["A1"] = "実施用 キャンペーン別 日予算（月予算÷日数、100円単位）"; ws4["A1"].font = Font(name=F, size=13, bold=True)
    ws4["A2"] = ("緑＝新設キャンペーン。3ストアは10月から 指名45% / ショッピング30% / カタログ・RTG25% の標準構成。"
                 "ショッピング・カタログはすべて通常価格品のみ（アウトレット除外）。イベント予備費は欄外。"); ws4["A2"].font = small
    hh = ["ストア", "ブランド/枠", "媒体", "キャンペーン"] + [f"{m} 日予算" for m in PLAN_MONTHS] + [f"{m} 月予算" for m in PLAN_MONTHS]
    for c_, x in enumerate(hh, 1):
        cell = ws4.cell(row=3, column=c_, value=x); cell.font = wb_; cell.fill = hdr; cell.alignment = center
    row = 4; nm = len(PLAN_MONTHS)
    for bi, (name, r9, r10, note) in enumerate(BRANDS):
        brow = first + bi
        for media, camp, share, new in CAMP[name]:
            ws4.cell(row=row, column=1, value="FULLMARKS"); ws4.cell(row=row, column=2, value=name)
            ws4.cell(row=row, column=3, value=media); ws4.cell(row=row, column=4, value=camp)
            for mi, m in enumerate(PLAN_MONTHS):
                mcol = 5 + nm + mi; dcol = 5 + mi
                ws4.cell(row=row, column=mcol, value=f"=ROUND(FULLMARKS内訳!{L(2+mi)}{brow}*{share},0)").number_format = YEN
                ws4.cell(row=row, column=dcol, value=f"=ROUND({L(mcol)}{row}/{days(m)},-2)").number_format = YEN
            if new:
                for c_ in range(1, 5 + 2 * nm): ws4.cell(row=row, column=c_).fill = newf
            row += 1
    for si, st in enumerate(STORES):
        scol = "DEF"[si]
        for media, camp, share in STORE_CAMP:
            ws4.cell(row=row, column=1, value=st); ws4.cell(row=row, column=2, value="ストア全体")
            ws4.cell(row=row, column=3, value=media); ws4.cell(row=row, column=4, value=camp)
            for mi, m in enumerate(PLAN_MONTHS):
                mcol = 5 + nm + mi; dcol = 5 + mi
                ws4.cell(row=row, column=mcol, value=f"=ROUND(月別予算!${scol}${month_rows[m]}*{share},0)").number_format = YEN
                ws4.cell(row=row, column=dcol, value=f"=ROUND({L(mcol)}{row}/{days(m)},-2)").number_format = YEN
            for c_ in range(1, 5 + 2 * nm): ws4.cell(row=row, column=c_).fill = newf
            row += 1
    tr4 = row
    ws4.cell(row=tr4, column=1, value="合計（通常運用）").font = bold
    for c_ in range(5, 5 + 2 * nm):
        cell = ws4.cell(row=tr4, column=c_, value=f"=SUM({L(c_)}4:{L(c_)}{tr4-1})"); cell.font = bold; cell.number_format = YEN; cell.fill = sub
    ws4.cell(row=tr4 + 1, column=1, value="イベント予備費（別枠）").font = bold
    for mi, m in enumerate(PLAN_MONTHS):
        cell = ws4.cell(row=tr4 + 1, column=5 + nm + mi, value=f"=月別予算!$I${month_rows[m]}"); cell.font = green; cell.number_format = YEN
    for rr in range(3, tr4 + 2):
        for c_ in range(1, 5 + 2 * nm): ws4.cell(row=rr, column=c_).border = border
    for col, w in zip("ABCD", (15, 26, 8, 40)): ws4.column_dimensions[col].width = w
    for c_ in range(5, 5 + 2 * nm): ws4.column_dimensions[L(c_)].width = 12
    ws4.freeze_panes = "E4"

    # ================= 9月実績 =================
    ws5 = wb.create_sheet("9月実績")
    ws5["A1"] = f"9月 実績（9/1〜9/{SEP_MTD_DAYS}）と月換算"; ws5["A1"].font = Font(name=F, size=13, bold=True)
    for c_, x in enumerate(["媒体", "キャンペーン", f"9/1〜9/{SEP_MTD_DAYS} 費用", "日割", "月換算(30日)"], 1):
        cell = ws5.cell(row=3, column=c_, value=x); cell.font = wb_; cell.fill = hdr; cell.alignment = center
    for i, (media, camp, amt) in enumerate(SEP_ACTUAL_BY_CAMP, 4):
        ws5.cell(row=i, column=1, value=media); ws5.cell(row=i, column=2, value=camp)
        cell = ws5.cell(row=i, column=3, value=amt); cell.font = blue; cell.number_format = YEN
        ws5.cell(row=i, column=4, value=f"=C{i}/{SEP_MTD_DAYS}").number_format = YEN
        ws5.cell(row=i, column=5, value=f"=D{i}*30").number_format = YEN
    e = 3 + len(SEP_ACTUAL_BY_CAMP) + 1
    ws5.cell(row=e, column=1, value="合計").font = bold
    for c_ in (3, 4, 5):
        cell = ws5.cell(row=e, column=c_, value=f"=SUM({L(c_)}4:{L(c_)}{e-1})"); cell.font = bold; cell.number_format = YEN; cell.fill = sub
    ws5.cell(row=e + 1, column=1, value="9月 通常運用計画").font = bold
    ws5.cell(row=e + 1, column=5, value=f"={K_SEP}").number_format = YEN
    ws5.cell(row=e + 2, column=1, value="差（月換算−計画）").font = bold
    ws5.cell(row=e + 2, column=5, value=f"=E{e}-E{e+1}").number_format = YEN
    ws5.cell(row=e + 3, column=1, value="注: 9/4夜〜9/6はMetaカタログ広告が停止していたため、その分は月換算が低めに出る。アウトレット枠は9/7に停止。").font = small
    for rr in range(3, e + 3):
        for c_ in range(1, 6): ws5.cell(row=rr, column=c_).border = border
    for col, w in zip("ABCDE", (12, 34, 16, 12, 14)): ws5.column_dimensions[col].width = w

    wb.calculation.fullCalcOnLoad = True
    wb.save(out)
    print("saved", out)


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else "reports/広告予算計画_2026-09改訂.xlsx")
